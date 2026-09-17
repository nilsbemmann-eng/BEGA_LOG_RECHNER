import io

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.models.party import Carrier
from app.models.tariff import Tariff, TariffRuleType, TariffStatus
from app.services.carrier_rate_import_service import CarrierRateMatrixParsingError, import_carrier_rate_matrix


def _stammdaten_workbook(carrier_rows: list[tuple]) -> bytes:
    """Baut ein synthetisches Workbook nach dem realen Stammdaten-Layout:
    Zeile 1-5 = Kopfzeilen (Land, Kfz-Zeichen, ISO Code 1, ISO Code 2, ISO Code
    3), ab Zeile 6 Spalte C = Frachtfuehrername, Spalte D+ = Preis je Land."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stammdaten"
    sheet.append(["Versendet", None, "Land", "Deutschland", "Polen", "Oesterreich"])
    sheet.append(["Bestaetigt", None, "Kfz-Zeichen", "D", "PL", "A"])
    sheet.append(["Preise ok", None, "ISO Code 1 (alpha-2)", "DE", "PL", "AT"])
    sheet.append(["STORNO", None, "ISO Code 2 (alpha-3)", "DEU", "POL", "AUT"])
    sheet.append(["Angemahnt", None, "ISO Code 3 (numeric)", 276, 616, 40])
    for row in carrier_rows:
        sheet.append([None, None, *row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_import_creates_carrier_and_tariff_with_valid_rates_only(db_session):
    content = _stammdaten_workbook([("ATB", 1.4, 1.55, None)])

    imported = import_carrier_rate_matrix(db_session, content)
    db_session.commit()

    assert imported == 1
    carrier = db_session.execute(select(Carrier).where(Carrier.name == "ATB")).scalars().one()
    tariff = db_session.execute(select(Tariff).where(Tariff.carrier_id == carrier.id)).scalars().one()
    assert tariff.status == TariffStatus.RELEASED
    rule = next(r for r in tariff.rules if r.rule_type == TariffRuleType.BASE_PLUS_KM)
    assert rule.parameters_json["price_per_km_by_country"] == {"DE": "1.4", "PL": "1.55"}


def test_import_skips_non_numeric_cells(db_session):
    """'keine' (kein Service), 'zu teuer', 'Sonderregelung' sind reale
    nicht-numerische Werte im Stammdaten-Blatt und duerfen nicht als Satz
    interpretiert werden (siehe docs/OFFENE_ENTSCHEIDUNGEN.md)."""
    content = _stammdaten_workbook([("BRW", "keine", "zu teuer", 1.2)])

    import_carrier_rate_matrix(db_session, content)
    db_session.commit()

    carrier = db_session.execute(select(Carrier).where(Carrier.name == "BRW")).scalars().one()
    tariff = db_session.execute(select(Tariff).where(Tariff.carrier_id == carrier.id)).scalars().one()
    rule = tariff.rules[0]
    assert rule.parameters_json["price_per_km_by_country"] == {"AT": "1.2"}


def test_import_skips_carrier_rows_without_any_valid_rate(db_session):
    content = _stammdaten_workbook([("ABATS", None, None, None)])

    imported = import_carrier_rate_matrix(db_session, content)
    db_session.commit()

    assert imported == 0
    assert db_session.execute(select(Carrier).where(Carrier.name == "ABATS")).scalars().first() is None


def test_reimport_archives_previous_tariff_and_creates_new_version(db_session):
    import_carrier_rate_matrix(db_session, _stammdaten_workbook([("ATB", 1.4, None, None)]))
    db_session.commit()

    import_carrier_rate_matrix(db_session, _stammdaten_workbook([("ATB", 1.5, None, None)]))
    db_session.commit()

    carrier = db_session.execute(select(Carrier).where(Carrier.name == "ATB")).scalars().one()
    tariffs = db_session.execute(select(Tariff).where(Tariff.carrier_id == carrier.id)).scalars().all()
    assert len(tariffs) == 2
    archived = next(t for t in tariffs if t.status == TariffStatus.ARCHIVED)
    released = next(t for t in tariffs if t.status == TariffStatus.RELEASED)
    assert archived.rules[0].parameters_json["price_per_km_by_country"] == {"DE": "1.4"}
    assert released.rules[0].parameters_json["price_per_km_by_country"] == {"DE": "1.5"}
    assert released.version == archived.version + 1


def test_reimport_preserves_manually_configured_parameters_not_in_stammdaten(db_session):
    """additional_unloading_point_price/toll_exempt/prefix_country_rate_overrides
    stammen nicht aus dem Stammdaten-Blatt (separate Formel-Ausnahmen im
    realen Excel) und duerfen bei einem Re-Import nicht verloren gehen."""
    import_carrier_rate_matrix(db_session, _stammdaten_workbook([("BABINSKI", 1.2, None, None)]))
    db_session.commit()

    carrier = db_session.execute(select(Carrier).where(Carrier.name == "BABINSKI")).scalars().one()
    tariff = db_session.execute(select(Tariff).where(Tariff.carrier_id == carrier.id)).scalars().one()
    tariff.rules[0].parameters_json = {
        **tariff.rules[0].parameters_json,
        "additional_unloading_point_price": "60.00",
        "toll_exempt": True,
    }
    db_session.commit()

    import_carrier_rate_matrix(db_session, _stammdaten_workbook([("BABINSKI", 1.3, None, None)]))
    db_session.commit()

    tariffs = db_session.execute(select(Tariff).where(Tariff.carrier_id == carrier.id)).scalars().all()
    released = next(t for t in tariffs if t.status == TariffStatus.RELEASED)
    params = released.rules[0].parameters_json
    assert params["price_per_km_by_country"] == {"DE": "1.3"}
    assert params["additional_unloading_point_price"] == "60.00"
    assert params["toll_exempt"] is True


def test_import_raises_on_missing_header(db_session):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Stammdaten"
    sheet.append(["Falsche", "Kopfzeile"])
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(CarrierRateMatrixParsingError):
        import_carrier_rate_matrix(db_session, buffer.getvalue())


def test_import_raises_when_stammdaten_sheet_missing(db_session):
    workbook = Workbook()
    workbook.active.title = "AndereTabelle"
    buffer = io.BytesIO()
    workbook.save(buffer)

    with pytest.raises(CarrierRateMatrixParsingError):
        import_carrier_rate_matrix(db_session, buffer.getvalue())
