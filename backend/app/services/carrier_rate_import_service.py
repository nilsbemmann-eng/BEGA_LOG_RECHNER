"""Import der Frachtfuehrer-Preistabelle ("Stammdaten"-Tabellenblatt aus
"Preise_2026_fuer_Wolke.xlsm", BEGA-Finetuning) - eine Matrix
Frachtfuehrer x Zielland -> EUR/km, die reale, manuell verhandelte
km-Preise fuer ca. 90 Subunternehmer enthaelt (siehe
docs/OFFENE_ENTSCHEIDUNGEN.md, app/tariff_engine/engine.py).

Nicht-numerische Zellen ("keine" = kein Service in dieses Land, "zu teuer",
"Sonderregelung", "?") werden bewusst uebersprungen, nicht als 0 interpretiert
- ein fehlender Satz fuehrt beim Audit zu `MissingCountryRateError` (manuelle
Pruefung) statt einer falschen Gratis-Fahrt.

Jeder Import legt fuer jeden betroffenen Frachtfuehrer eine NEUE Tarifversion
an (Tarife sind unveraenderlich, sobald sie in einem `AuditResult` referenziert
wurden, siehe `app/models/tariff.py`) und uebernimmt dabei
`additional_unloading_point_price`/`toll_exempt`/`prefix_country_rate_overrides`
aus der bisherigen Tarifversion unveraendert, da diese Parameter nicht aus
der Stammdaten-Tabelle stammen (sondern aus separaten Formel-Ausnahmen im
Excel, siehe docs/OFFENE_ENTSCHEIDUNGEN.md).
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.party import Carrier
from app.models.tariff import Tariff, TariffRule, TariffRuleType, TariffStatus

_SHEET_NAME = "Stammdaten"
_HEADER_LABEL = "ISO Code 1 (alpha-2)"
_HEADER_LABEL_COLUMN_INDEX = 2  # Spalte C (0-basiert)
_CARRIER_NAME_COLUMN_INDEX = 2  # ebenfalls Spalte C, ab der Zeile nach dem Header

# Spalte-C-Kopfzeilen-Labels, die selbst KEIN Frachtfuehrername sind (siehe
# reale Vorlage: "ISO Code 3 (numeric)" steht z. B. in derselben Spalte wie
# die Frachtfuehrernamen und enthaelt in den Laenderspalten numerische
# ISO-Nummerncodes - ohne diese Ausschlussliste wuerde die Kopfzeile selbst
# faelschlich als Frachtfuehrer mit gueltigen "km-Preisen" importiert).
_NON_CARRIER_LABELS = {
    "Land",
    "Kfz-Zeichen",
    "ISO Code 1 (alpha-2)",
    "ISO Code 2 (alpha-3)",
    "ISO Code 3 (numeric)",
}


class CarrierRateMatrixParsingError(ValueError):
    """Wird geworfen, wenn die Excel-Struktur nicht der erwarteten Stammdaten-Vorlage entspricht."""


def _slugify_carrier_code(name: str) -> str:
    normalized = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode("ascii")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", normalized).strip("-").upper()
    return slug or "CARRIER"


def _find_or_create_carrier(db: Session, name: str) -> Carrier:
    carrier = db.execute(select(Carrier).where(Carrier.name == name)).scalars().first()
    if carrier is not None:
        return carrier

    base_code = _slugify_carrier_code(name)
    code = base_code
    suffix = 1
    while db.execute(select(Carrier).where(Carrier.carrier_code == code)).scalars().first() is not None:
        suffix += 1
        code = f"{base_code}-{suffix}"

    carrier = Carrier(name=name, carrier_code=code)
    db.add(carrier)
    db.flush()
    return carrier


def _upsert_tariff(db: Session, carrier: Carrier, price_per_km_by_country: dict[str, str]) -> None:
    existing = db.execute(
        select(Tariff).where(Tariff.carrier_id == carrier.id, Tariff.status == TariffStatus.RELEASED)
    ).scalars().all()

    previous_params: dict = {}
    next_version = 1
    for tariff in existing:
        rule = next((r for r in tariff.rules if r.rule_type == TariffRuleType.BASE_PLUS_KM), None)
        if rule is not None:
            previous_params = rule.parameters_json
        tariff.status = TariffStatus.ARCHIVED
        next_version = max(next_version, tariff.version + 1)

    new_params = {
        **{k: v for k, v in previous_params.items() if k != "price_per_km_by_country"},
        "base_price": previous_params.get("base_price", "0"),
        "minimum_km": previous_params.get("minimum_km", "0"),
        "price_per_km_by_country": price_per_km_by_country,
    }

    tariff = Tariff(
        tariff_code=f"{carrier.carrier_code}-{next_version}",
        name=f"{carrier.name} (Stammdaten-Import)",
        carrier_id=carrier.id,
        valid_from=date.today(),
        valid_to=None,
        status=TariffStatus.RELEASED,
        version=next_version,
        rules=[TariffRule(rule_type=TariffRuleType.BASE_PLUS_KM, parameters_json=new_params, priority=0)],
    )
    db.add(tariff)


def import_carrier_rate_matrix(db: Session, file_bytes: bytes) -> int:
    workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    if _SHEET_NAME not in workbook.sheetnames:
        raise CarrierRateMatrixParsingError(f"Tabellenblatt '{_SHEET_NAME}' nicht gefunden")
    sheet = workbook[_SHEET_NAME]
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    header_row_index = next(
        (i for i, row in enumerate(rows) if len(row) > _HEADER_LABEL_COLUMN_INDEX and row[_HEADER_LABEL_COLUMN_INDEX] == _HEADER_LABEL),
        None,
    )
    if header_row_index is None:
        raise CarrierRateMatrixParsingError(f"Kopfzeile '{_HEADER_LABEL}' nicht in Spalte C gefunden")

    header_row = rows[header_row_index]
    country_codes = header_row[_HEADER_LABEL_COLUMN_INDEX + 1 :]

    imported = 0
    for row in rows[header_row_index + 1 :]:
        if len(row) <= _CARRIER_NAME_COLUMN_INDEX or not row[_CARRIER_NAME_COLUMN_INDEX]:
            continue
        carrier_name = str(row[_CARRIER_NAME_COLUMN_INDEX]).strip()
        if carrier_name in _NON_CARRIER_LABELS:
            continue
        rate_cells = row[_HEADER_LABEL_COLUMN_INDEX + 1 :]

        price_per_km_by_country: dict[str, str] = {}
        for country_code, value in zip(country_codes, rate_cells):
            if not country_code or not isinstance(value, (int, float)):
                continue  # "keine"/"zu teuer"/"Sonderregelung"/"?"/leer werden bewusst uebersprungen
            price_per_km_by_country[str(country_code)] = str(value)

        if not price_per_km_by_country:
            continue

        carrier = _find_or_create_carrier(db, carrier_name)
        _upsert_tariff(db, carrier, price_per_km_by_country)
        imported += 1

    db.flush()
    return imported
