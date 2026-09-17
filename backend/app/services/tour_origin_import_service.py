"""Import der "Gebietsrelationen"-Stammdatentabelle (BEGA-Finetuning,
Nutzervorgabe) in die Absender-Matrix (`TourOriginMapping`).

Erwartete Spalten (siehe reale Vorlage "Gebietsrelationen.xlsx"):
Matchcode | Bezeichnung | Praefix | Absender | Absenderadresse

`Absenderadresse` hat das Format "<Laendercode> <PLZ> <Ort> <Strasse>", z. B.
"PL 39-300 Mielec ul. Wojska Polskiego 3" oder "D 21129 Hamburg Am Ballinkai 1".
Diese Tabelle ist die vollstaendige, aktuelle Referenz (kein inkrementelles
Update) - ein Import ersetzt daher den gesamten bisherigen *aktuellen* Inhalt
der Absender-Matrix, analog zu einem periodischen Stammdatenabgleich. Anders
als frueher werden bestehende Zeilen dabei NICHT geloescht, sondern
versioniert (siehe `app/models/tour_origin_mapping.py`): jede (Praefix,
Matchcode)-Kombination aus der neuen Datei wird als neue Version angelegt
(oder bleibt unveraendert `is_current`, falls sie schon existiert und
identisch ist), alle zuvor aktuellen Zeilen, die im neuen Import nicht mehr
vorkommen, werden auf `is_current=False` gesetzt - so bleibt nachvollziehbar,
welche Absenderadresse zu welchem Zeitpunkt fuer einen Praefix galt.
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.address import Address
from app.models.tour_origin_mapping import TourOriginMapping

_EXPECTED_HEADER = ["Matchcode", "Bezeichnung", "Präfix", "Absender", "Absenderadresse"]
_ADDRESS_RE = re.compile(r"^([A-Z]{1,2})\s+(\S+)\s+(\S+)\s*(.*)$")
_COUNTRY_PREFIX_MAP = {"D": "DE"}


class TourOriginMatrixParsingError(ValueError):
    """Wird geworfen, wenn die Excel-Struktur nicht der erwarteten Vorlage entspricht."""


def _parse_absenderadresse(raw: str) -> Address:
    match = _ADDRESS_RE.match(raw.strip())
    if not match:
        raise TourOriginMatrixParsingError(f"Absenderadresse konnte nicht geparst werden: {raw!r}")
    country_prefix, postal_code, city, street = match.groups()
    return Address(
        original_text=raw,
        street=street or None,
        postal_code=postal_code,
        city=city,
        country_code=_COUNTRY_PREFIX_MAP.get(country_prefix.upper(), country_prefix.upper()),
    )


def _address_identity(address: Address | None) -> tuple:
    if address is None:
        return (None, None, None, None)
    return (address.street, address.postal_code, address.city, address.country_code)


def import_tour_origin_matrix(db: Session, file_bytes: bytes) -> int:
    workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    sheet = workbook.active
    rows = [list(row) for row in sheet.iter_rows(values_only=True)]
    if not rows:
        raise TourOriginMatrixParsingError("Datei enthaelt keine Zeilen")

    header = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    if header[: len(_EXPECTED_HEADER)] != _EXPECTED_HEADER:
        raise TourOriginMatrixParsingError(
            f"Unerwartete Spaltenkopfzeile {header} - erwartet: {_EXPECTED_HEADER}"
        )

    current_rows = db.execute(
        select(TourOriginMapping).where(TourOriginMapping.is_current.is_(True))
    ).scalars().all()
    current_by_key: dict[tuple[str, str], TourOriginMapping] = {
        (m.tour_number_prefix, m.matchcode): m for m in current_rows
    }
    seen_keys: set[tuple[str, str]] = set()

    imported = 0
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        matchcode, description, prefix, _absender, absenderadresse = (list(row) + [None] * 5)[:5]

        prefix = str(prefix).strip()
        matchcode = str(matchcode).strip()
        description = str(description).strip() if description else None
        origin_address = _parse_absenderadresse(str(absenderadresse)) if absenderadresse else None

        key = (prefix, matchcode)
        seen_keys.add(key)
        existing = current_by_key.get(key)

        if existing is not None and existing.description == description and _address_identity(existing.origin_address) == _address_identity(origin_address):
            imported += 1
            continue  # unveraendert - keine neue Version noetig

        if existing is not None:
            existing.is_current = False

        db.add(
            TourOriginMapping(
                tour_number_prefix=prefix,
                matchcode=matchcode,
                description=description,
                origin_address=origin_address,
                version=(existing.version + 1) if existing is not None else 1,
                is_current=True,
            )
        )
        imported += 1

    # Zeilen, die es in der neuen Datei nicht mehr gibt, sind ab jetzt nicht
    # mehr aktuell (aber bleiben als Historie erhalten).
    for key, mapping in current_by_key.items():
        if key not in seen_keys:
            mapping.is_current = False

    db.flush()
    return imported
