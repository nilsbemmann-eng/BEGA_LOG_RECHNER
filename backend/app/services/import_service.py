"""Ladelisten-Import (Abschnitt 4.5, 16.2).

Unterstuetzt CSV und XLSX. Unbekannte Spalten gehen nicht verloren, sondern
werden je Zeile gesammelt und als JSON-Bericht referenziert
(`ImportJob.error_report_reference`) - so bleiben sie im Importprotokoll
erhalten (Akzeptanzkriterium 16.2), ohne das Datenmodell der Sendung mit
Ad-hoc-Feldern zu belasten.
"""
from __future__ import annotations

import csv
import io
import json
import os
from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.address import Address
from app.models.import_job import ImportJob, ImportJobStatus, ImportSourceType
from app.models.shipment import Shipment
from app.normalization.columns import map_columns
from app.normalization.numbers import NumberParsingError, parse_german_decimal

_NUMERIC_FIELDS = {"weight_kg", "pallets", "loading_meters", "volume_m3", "freight_amount", "surcharge_amount"}


def _read_rows(file_bytes: bytes, filename: str) -> tuple[list[str], list[list[str]]]:
    if filename.lower().endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        reader = csv.reader(io.StringIO(text), delimiter=";")
        rows = list(reader)
        if not rows:
            return [], []
        return rows[0], rows[1:]

    if filename.lower().endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [[("" if cell is None else str(cell)) for cell in row] for row in sheet.iter_rows(values_only=True)]
        if not rows:
            return [], []
        return rows[0], rows[1:]

    raise ValueError(f"Nicht unterstuetztes Ladeliste-Format: {filename}")


def import_ladeliste(db: Session, file_bytes: bytes, filename: str, report_storage_path: str) -> ImportJob:
    header, data_rows = _read_rows(file_bytes, filename)
    mapping = map_columns(header)

    job = ImportJob(
        source_type=ImportSourceType.LADELISTE_UPLOAD,
        started_at=datetime.now(timezone.utc),
        status=ImportJobStatus.RUNNING,
        records_total=len(data_rows),
    )
    db.add(job)
    db.flush()

    unmapped_report: list[dict] = []
    successful = 0
    failed = 0

    for row_index, row in enumerate(data_rows, start=1):
        row_dict = dict(zip(header, row, strict=False))
        try:
            shipment = _row_to_shipment(row_dict, mapping.field_to_column)
            db.add(shipment)
            successful += 1
        except (NumberParsingError, ValueError) as exc:
            failed += 1
            unmapped_report.append({"row": row_index, "error": str(exc), "raw_row": row_dict})
            continue

        if mapping.unmapped_columns:
            unmapped_report.append(
                {
                    "row": row_index,
                    "shipment_number": row_dict.get(mapping.field_to_column.get("shipment_number", ""), None),
                    "unmapped_fields": {column: row_dict.get(column) for column in mapping.unmapped_columns},
                }
            )

    job.completed_at = datetime.now(timezone.utc)
    job.records_successful = successful
    job.records_failed = failed
    job.status = ImportJobStatus.COMPLETED if failed == 0 else ImportJobStatus.COMPLETED_WITH_ERRORS

    if unmapped_report:
        os.makedirs(report_storage_path, exist_ok=True)
        report_path = os.path.join(report_storage_path, f"{job.id}.json")
        with open(report_path, "w", encoding="utf-8") as handle:
            json.dump(unmapped_report, handle, ensure_ascii=False, indent=2, default=str)
        job.error_report_reference = report_path

    db.flush()
    return job


def _row_to_shipment(row: dict, field_to_column: dict[str, str]) -> Shipment:
    def get_raw(field: str) -> str | None:
        column = field_to_column.get(field)
        if not column:
            return None
        value = row.get(column)
        return value.strip() if isinstance(value, str) and value.strip() else None

    def get_decimal(field: str) -> Decimal | None:
        raw = get_raw(field)
        return parse_german_decimal(raw) if raw else None

    origin_text = get_raw("origin_address")
    destination_text = get_raw("destination_address")
    freight_amount = get_decimal("freight_amount")
    pallets_raw = get_decimal("pallets")

    return Shipment(
        shipment_number=get_raw("shipment_number"),
        origin_address=Address(original_text=origin_text) if origin_text else None,
        destination_address=Address(original_text=destination_text) if destination_text else None,
        weight_kg=get_decimal("weight_kg"),
        pallets=int(pallets_raw) if pallets_raw is not None else None,
        loading_meters=get_decimal("loading_meters"),
        volume_m3=get_decimal("volume_m3"),
        invoice_amount=freight_amount,
    )
